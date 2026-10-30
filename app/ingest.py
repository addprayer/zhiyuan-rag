"""Document extraction, chunking and indexing."""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

from .chunking import Segment, chunk_segments
from .config import AppConfig
from .db import Database, new_id, now_iso
from .embeddings import build_embedder, signature_for
from .vector_store import pack_vector

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}
MAX_FILE_SIZE = 20 * 1024 * 1024


class IngestError(RuntimeError):
    pass


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise IngestError("无法识别文本编码，请另存为 UTF-8 后重试")


def _markdown_segments(text: str) -> list[Segment]:
    segments: list[Segment] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        content = "\n".join(buffer).strip()
        if content:
            segments.append(Segment(content, page=None, section=section))
        buffer = []

    for line in text.splitlines():
        match = re.match(r"^#{1,6}\s+(.+?)\s*#*\s*$", line)
        if match:
            flush()
            section = match.group(1).strip()
        else:
            buffer.append(line)
    flush()
    return segments


def _pdf_segments(path: Path) -> tuple[list[Segment], int]:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:  # PyPDF2 is a compatible fallback for older installs
        from PyPDF2 import PdfReader  # type: ignore
    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise IngestError(f"PDF 读取失败：{exc}") from exc
    segments: list[Segment] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            segments.append(Segment(text, page=index))
    return segments, len(reader.pages)


def _docx_segments(path: Path) -> list[Segment]:
    try:
        from docx import Document
    except ImportError as exc:
        raise IngestError("未安装 python-docx，无法解析 .docx") from exc
    try:
        document = Document(str(path))
    except Exception as exc:
        raise IngestError(f"Word 读取失败：{exc}") from exc

    segments: list[Segment] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        content = "\n\n".join(part for part in buffer if part.strip()).strip()
        if content:
            segments.append(Segment(content, page=None, section=section))
        buffer = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        if style.lower().startswith("heading") or style.startswith("标题"):
            flush()
            section = text
        else:
            buffer.append(text)
    flush()

    table_lines: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_lines.append(" | ".join(cells))
    if table_lines:
        segments.append(Segment("\n".join(table_lines), page=None, section="表格"))
    return segments


def extract_segments(path: Path) -> tuple[list[Segment], int | None]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        text = _read_text(path)
        if suffix in {".md", ".markdown"}:
            return _markdown_segments(text), None
        return [Segment(text)], None
    if suffix == ".pdf":
        return _pdf_segments(path)
    if suffix == ".docx":
        return _docx_segments(path), None
    raise IngestError(f"不支持的文件类型：{suffix}")


def ingest_file(
    db: Database,
    settings: dict[str, Any],
    kb: Any,
    path: Path,
    display_name: str,
    stored_path: str = "",
) -> dict[str, Any]:
    signature = signature_for(settings)
    if kb["embedding_signature"] != signature:
        raise IngestError(
            "当前 Embedding 配置与知识库建立时不一致，请先在设置中重建该知识库索引"
        )

    segments, page_count = extract_segments(path)
    chunks = chunk_segments(
        segments,
        chunk_size=int(settings.get("chunk_size", 520)),
        overlap=int(settings.get("chunk_overlap", 80)),
    )
    if not chunks:
        raise IngestError("没有从文档中提取到可索引的文字")

    embedder = build_embedder(settings)
    embed_texts = [
        f"{chunk.section}\n{chunk.content}".strip() if chunk.section else chunk.content
        for chunk in chunks
    ]
    vectors = embedder.embed_documents(embed_texts)
    if len(vectors) != len(chunks):
        raise IngestError("向量数量与文本片段数量不一致")

    doc_id = new_id("doc")
    created_at = now_iso()
    with db.conn() as conn:
        conn.execute(
            """
            INSERT INTO documents
                (id, kb_id, name, stored_path, file_type, size, page_count,
                 chunk_count, status, error, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'ready', NULL, ?)
            """,
            (
                doc_id,
                kb["id"],
                display_name,
                stored_path,
                path.suffix.lower().lstrip("."),
                path.stat().st_size,
                page_count,
                len(chunks),
                created_at,
            ),
        )
        conn.executemany(
            """
            INSERT INTO chunks
                (doc_id, kb_id, chunk_index, content, page, section,
                 embedding, dim, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    doc_id,
                    kb["id"],
                    chunk.index,
                    chunk.content,
                    chunk.page,
                    chunk.section,
                    pack_vector(vector),
                    int(len(vector)),
                    created_at,
                )
                for chunk, vector in zip(chunks, vectors)
            ],
        )
    return {
        "id": doc_id,
        "kb_id": kb["id"],
        "name": display_name,
        "file_type": path.suffix.lower().lstrip("."),
        "size": path.stat().st_size,
        "page_count": page_count,
        "chunk_count": len(chunks),
        "status": "ready",
        "created_at": created_at,
    }


def reindex_kb(db: Database, settings: dict[str, Any], kb_id: str) -> int:
    kb = db.one("SELECT * FROM kbs WHERE id = ?", (kb_id,))
    if not kb:
        raise IngestError("知识库不存在")
    embedder = build_embedder(settings)
    docs = db.all("SELECT id FROM documents WHERE kb_id = ?", (kb_id,))
    total = 0
    for doc in docs:
        rows = db.all(
            "SELECT id, content, section FROM chunks WHERE doc_id = ? ORDER BY chunk_index",
            (doc["id"],),
        )
        if not rows:
            continue
        texts = [
            f"{row['section']}\n{row['content']}".strip()
            if row["section"]
            else str(row["content"])
            for row in rows
        ]
        vectors = embedder.embed_documents(texts)
        with db.conn() as conn:
            for row, vector in zip(rows, vectors):
                conn.execute(
                    "UPDATE chunks SET embedding = ?, dim = ? WHERE id = ?",
                    (pack_vector(vector), int(len(vector)), row["id"]),
                )
        total += len(rows)
    db.execute(
        "UPDATE kbs SET embedding_signature = ? WHERE id = ?",
        (signature_for(settings), kb_id),
    )
    return total


def document_dict(row: Any) -> dict[str, Any]:
    return {
        "id": row["id"],
        "kb_id": row["kb_id"],
        "name": row["name"],
        "file_type": row["file_type"],
        "size": row["size"],
        "page_count": row["page_count"],
        "chunk_count": row["chunk_count"],
        "status": row["status"],
        "error": row["error"],
        "created_at": row["created_at"],
    }


def seed_demo(db: Database, config: AppConfig, store: Any) -> bool:
    existing = db.one("SELECT COUNT(*) AS count FROM kbs")
    if existing and int(existing["count"]) > 0:
        return False
    settings = store.get()
    signature = signature_for(settings)
    kb_id = new_id("kb")
    db.execute(
        "INSERT INTO kbs (id, name, description, embedding_signature, created_at) "
        "VALUES (?, ?, ?, ?, ?)",
        (
            kb_id,
            "云栖科技 · 员工知识库",
            "演示数据：员工手册、信息安全规范、客户支持流程",
            signature,
            now_iso(),
        ),
    )
    kb = db.one("SELECT * FROM kbs WHERE id = ?", (kb_id,))
    if not config.samples_dir.exists():
        return True
    for path in sorted(config.samples_dir.glob("*.md")):
        try:
            ingest_file(
                db,
                settings,
                kb,
                path,
                path.name,
                stored_path=f"samples/{path.name}",
            )
        except Exception:
            logger.exception("Failed to seed sample %s", path.name)
    return True
